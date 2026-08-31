# -*- coding: utf-8 -*-
"""Normalize the two school club workbooks into the runtime seed catalog."""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


GRADE_NAMES = {
    "一": "一年级",
    "二": "二年级",
    "三": "三年级",
    "四": "四年级",
    "五": "五年级",
}
SHEET_RULES = {
    "东校区社团": {"site": "东校区", "info_only_from": 43},
    "西校区社团": {"site": "西校区", "info_only_from": 29},
}


def text(value):
    return str(value or "").strip()


def grades(value):
    raw = text(value)
    return [name for numeral, name in GRADE_NAMES.items() if numeral in raw]


def capacity(value):
    if isinstance(value, (int, float)):
        return int(value)
    numbers = [int(item) for item in re.findall(r"\d+", text(value))]
    return sum(numbers)


def normalize_workbook(path):
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    rule = SHEET_RULES.get(worksheet.title)
    if not rule:
        raise ValueError(f"不支持的工作表：{worksheet.title}")

    courses = []
    for row_number in range(3, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, 12)]
        source_number = values[0]
        name = text(values[3])
        if not source_number or not name:
            continue

        teacher = text(values[1])
        weekday = text(values[5])
        location = text(values[6])
        note = text(values[7])
        quota = capacity(values[9])
        is_tail_information = int(source_number) >= rule["info_only_from"]
        required_complete = bool(teacher and weekday and location)
        if quota > 0:
            selection_mode = "selectable"
        elif is_tail_information and required_complete:
            selection_mode = "info_only"
        else:
            selection_mode = "draft"

        eligible_source = values[8] if selection_mode == "selectable" else values[4]
        eligible_grades = grades(eligible_source)
        if not eligible_grades:
            selection_mode = "draft"

        site_key = "east" if rule["site"] == "东校区" else "west"
        courses.append({
            "id": f"benbu-2026s1-{site_key}-{int(source_number):03d}",
            "campus": rule["site"],
            "name": name,
            "teacher": teacher,
            "weekday": weekday,
            "location": location,
            # Legacy schema requires a positive capacity. Non-selectable rows never expose or use it.
            "capacity": quota or 1,
            "note": note,
            "selection_mode": selection_mode,
            "eligible_grades": eligible_grades,
            "source_number": int(source_number),
        })
    return courses


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    courses = []
    for workbook in args.workbooks:
        courses.extend(normalize_workbook(workbook))
    ids = [course["id"] for course in courses]
    if len(ids) != len(set(ids)):
        raise ValueError("课程编号重复")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(courses, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    summary = {
        mode: sum(course["selection_mode"] == mode for course in courses)
        for mode in ("selectable", "info_only", "draft")
    }
    print(json.dumps({"total": len(courses), **summary}, ensure_ascii=False))


if __name__ == "__main__":
    main()
