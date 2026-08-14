"""食堂周菜单 Excel 解析器。

解析依据是表内语义标签，不依赖工作表名称、固定行数或文件名。
Excel 是结构化主数据；图片仅用于展示和人工核对。
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


WEEKDAY_LABELS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
NON_SERVICE_WORDS = ("放假", "停餐", "不供餐", "秋游", "春游", "外出", "调休")


class MenuWorkbookError(ValueError):
    """上传文件不是可识别的食堂菜单。"""


def _text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _merged_value(ws, row: int, column: int):
    cell = ws.cell(row, column)
    if cell.value not in (None, ""):
        return cell.value
    coordinate = cell.coordinate
    for merged in ws.merged_cells.ranges:
        if coordinate in merged:
            return ws.cell(merged.min_row, merged.min_col).value
    return None


def _parse_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _number(value, unit_pattern=""):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "")
    if unit_pattern:
        text = re.sub(unit_pattern, "", text, flags=re.I)
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def _percent(value, number_format: str):
    number = _number(value)
    if number is None:
        return None
    if "%" in (number_format or "") or number <= 1.5:
        number *= 100
    return round(number, 2)


def _unique(values):
    result = []
    for value in values:
        text = _text(value)
        if text and text not in result:
            result.append(text)
    return result


def _sheet_score(ws) -> int:
    score = 0
    for row in ws.iter_rows():
        for cell in row:
            text = _text(cell.value)
            if text in ("A套餐", "B套餐"):
                score += 3
            elif text in ("膳食品类", "热量", "热量（kcal）"):
                score += 2
    return score


def _issue(severity, code, message, plan_date=None, field=None):
    result = {"severity": severity, "code": code, "message": message}
    if plan_date:
        result["plan_date"] = plan_date
    if field:
        result["field"] = field
    return result


def parse_menu_workbook(content: bytes) -> dict:
    """把食堂周菜单解析为按日期组织的 A/B 餐结构。"""
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise MenuWorkbookError("Excel 文件无法打开，请确认文件未损坏且格式为 .xlsx") from exc

    candidates = sorted(workbook.worksheets, key=_sheet_score, reverse=True)
    if not candidates or _sheet_score(candidates[0]) < 4:
        raise MenuWorkbookError("没有找到包含 A套餐、B套餐和膳食品类的菜单工作表")
    ws = candidates[0]

    meal_header_row = None
    for row in range(1, min(ws.max_row, 40) + 1):
        labels = [_text(_merged_value(ws, row, col)) for col in range(1, ws.max_column + 1)]
        if "A套餐" in labels and "B套餐" in labels:
            meal_header_row = row
            break
    if not meal_header_row or meal_header_row == 1:
        raise MenuWorkbookError("没有找到日期与 A/B 套餐表头")

    date_row = meal_header_row - 1
    title = _text(ws["A1"].value)
    row_labels = {
        row: _text(_merged_value(ws, row, 2))
        for row in range(meal_header_row + 1, ws.max_row + 1)
    }
    nutrition_start = min(
        (row for row, label in row_labels.items() if "热量" in label),
        default=ws.max_row + 1,
    )
    name_rows = [
        row for row, label in row_labels.items()
        if row < nutrition_start and "品名" in label and label != "膳食品类"
    ]
    ingredient_rows = [
        row for row, label in row_labels.items()
        if row < nutrition_start and "份量" in label
    ]

    nutrient_rows = {}
    for row, label in row_labels.items():
        if "热量" in label:
            nutrient_rows["calories_kcal"] = row
        elif "蛋白质" in label:
            nutrient_rows["protein_pct"] = row
        elif "脂肪" in label:
            nutrient_rows["fat_pct"] = row
        elif "维生素" in label.lower() and "c" in label.lower():
            nutrient_rows["vitamin_c_mg"] = row

    issues = []
    days = []
    seen_dates = set()
    for column in range(1, ws.max_column + 1):
        parsed_date = _parse_date(_merged_value(ws, date_row, column), workbook.epoch)
        if not parsed_date:
            continue
        plan_date = parsed_date.isoformat()
        if plan_date in seen_dates:
            continue
        seen_dates.add(plan_date)
        if column >= ws.max_column:
            issues.append(_issue("error", "missing_meal_column", f"{plan_date} 缺少 B 餐列", plan_date))
            continue

        actual_weekday = parsed_date.weekday() + 1
        stated_weekday = _text(_merged_value(ws, date_row, column + 1))
        expected_weekday = WEEKDAY_LABELS[actual_weekday - 1]
        if stated_weekday and stated_weekday != expected_weekday:
            issues.append(_issue(
                "error", "weekday_mismatch",
                f"{plan_date} 标注为{stated_weekday}，实际是{expected_weekday}", plan_date, "weekday",
            ))

        type_a = _text(_merged_value(ws, meal_header_row, column))
        type_b = _text(_merged_value(ws, meal_header_row, column + 1))
        pair_values = [
            _text(_merged_value(ws, row, meal_column))
            for row in range(meal_header_row, nutrition_start)
            for meal_column in (column, column + 1)
        ]
        service_note = next(
            (value for value in pair_values if any(word in value for word in NON_SERVICE_WORDS)),
            "",
        )
        service_status = "normal" if type_a == "A套餐" and type_b == "B套餐" else "no_service"
        day = {
            "plan_date": plan_date,
            "weekday": actual_weekday,
            "weekday_label": expected_weekday,
            "service_status": service_status,
            "service_note": service_note or ("非供餐日" if service_status != "normal" else ""),
            "meals": [],
        }
        if service_status != "normal":
            days.append(day)
            continue

        for plan_type, meal_column in (("A", column), ("B", column + 1)):
            menu_items = _unique(_merged_value(ws, row, meal_column) for row in name_rows)
            ingredients = _unique(_merged_value(ws, row, meal_column) for row in ingredient_rows)
            if not menu_items:
                issues.append(_issue(
                    "error", "missing_meal_items",
                    f"{plan_date} {plan_type}餐没有识别到菜品", plan_date, f"meal_{plan_type}",
                ))
            meal = {
                "plan_type": plan_type,
                "plan_name": "、".join(menu_items),
                "menu_items": menu_items,
                "ingredients": ingredients,
                "calories_kcal": None,
                "protein_pct": None,
                "fat_pct": None,
                "vitamin_c_mg": None,
            }
            if "calories_kcal" in nutrient_rows:
                meal["calories_kcal"] = _number(
                    _merged_value(ws, nutrient_rows["calories_kcal"], meal_column), r"kcal",
                )
            for field in ("protein_pct", "fat_pct"):
                if field in nutrient_rows:
                    row = nutrient_rows[field]
                    meal[field] = _percent(
                        _merged_value(ws, row, meal_column),
                        ws.cell(row, meal_column).number_format,
                    )
            if "vitamin_c_mg" in nutrient_rows:
                meal["vitamin_c_mg"] = _number(
                    _merged_value(ws, nutrient_rows["vitamin_c_mg"], meal_column), r"mg",
                )

            if meal["calories_kcal"] is None:
                issues.append(_issue(
                    "warning", "missing_calories",
                    f"{plan_date} {plan_type}餐缺少热量", plan_date, f"meal_{plan_type}.calories_kcal",
                ))
            elif not 200 <= meal["calories_kcal"] <= 1500:
                issues.append(_issue(
                    "warning", "calories_outlier",
                    f"{plan_date} {plan_type}餐热量 {meal['calories_kcal']:g} kcal 需要确认",
                    plan_date, f"meal_{plan_type}.calories_kcal",
                ))
            for field, label in (("protein_pct", "蛋白质"), ("fat_pct", "脂肪")):
                value = meal[field]
                if value is not None and not 0 <= value <= 100:
                    issues.append(_issue(
                        "error", "percentage_outlier",
                        f"{plan_date} {plan_type}餐{label}为 {value:g}%，请修正Excel后重新上传",
                        plan_date, f"meal_{plan_type}.{field}",
                    ))
            day["meals"].append(meal)
        days.append(day)

    if not days:
        raise MenuWorkbookError("没有识别到有效日期列")
    days.sort(key=lambda item: item["plan_date"])
    date_span = (date.fromisoformat(days[-1]["plan_date"]) - date.fromisoformat(days[0]["plan_date"])).days
    normal_days = [day for day in days if day["service_status"] == "normal"]
    if date_span > 6 or len(normal_days) > 7:
        issues.append(_issue(
            "error",
            "multiple_weeks_detected",
            "一个周次只能发布一个自然周，请将跨周Excel拆分后分别上传",
            field="date_range",
        ))
    return {
        "title": title,
        "sheet_name": ws.title,
        "date_start": days[0]["plan_date"],
        "date_end": days[-1]["plan_date"],
        "days": days,
        "issues": issues,
        "service_day_count": len(normal_days),
    }
