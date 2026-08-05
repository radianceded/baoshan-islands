import io
import json
import os
import sqlite3
import tempfile
import unittest
from datetime import date, timedelta

from openpyxl import Workbook

from nutrition import models as nutrition_models
from server import app as server_app
from server.menu_excel import parse_menu_workbook


def build_menu_workbook(start=date(2026, 8, 3), no_service_weekday=None, fat_value=0.28):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生菜单"
    sheet["A1"] = "宝山实验小学学生菜单"
    sheet.cell(4, 2, "膳食品类")
    rows = {
        5: "主食品名",
        6: "主食份量",
        7: "副食品名",
        8: "副食份量",
        9: "热量（kcal）",
        10: "蛋白质（%）",
        11: "脂肪（%）",
        12: "维生素C（mg）",
    }
    for row, label in rows.items():
        sheet.cell(row, 2, label)

    for index in range(5):
        plan_date = start + timedelta(days=index)
        column = 3 + index * 2
        sheet.cell(3, column, plan_date)
        sheet.cell(3, column + 1, f"周{'一二三四五'[index]}")
        if index + 1 == no_service_weekday:
            sheet.cell(4, column, "学校秋游，不供餐")
            continue
        sheet.cell(4, column, "A套餐")
        sheet.cell(4, column + 1, "B套餐")
        for offset, plan_type in enumerate(("A", "B")):
            meal_column = column + offset
            sheet.cell(5, meal_column, f"米饭{plan_type}")
            sheet.cell(6, meal_column, "大米100克")
            sheet.cell(7, meal_column, f"鸡肉时蔬{plan_type}")
            sheet.cell(8, meal_column, "鸡肉、青菜")
            sheet.cell(9, meal_column, 520 + offset * 20)
            sheet.cell(10, meal_column, 0.24 + offset * 0.01)
            sheet.cell(10, meal_column).number_format = "0%"
            sheet.cell(11, meal_column, fat_value)
            sheet.cell(11, meal_column).number_format = "0%"
            sheet.cell(12, meal_column, 32 + offset)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class MenuWorkbookParserTests(unittest.TestCase):
    def test_parses_dates_meals_and_nutrition_by_semantic_labels(self):
        parsed = parse_menu_workbook(build_menu_workbook(no_service_weekday=3))

        self.assertEqual(parsed["date_start"], "2026-08-03")
        self.assertEqual(parsed["date_end"], "2026-08-07")
        self.assertEqual(parsed["service_day_count"], 4)
        self.assertEqual(parsed["days"][2]["service_status"], "no_service")
        monday_a = parsed["days"][0]["meals"][0]
        self.assertEqual(monday_a["plan_name"], "米饭A、鸡肉时蔬A")
        self.assertEqual(monday_a["calories_kcal"], 520)
        self.assertEqual(monday_a["protein_pct"], 24)
        self.assertEqual(monday_a["fat_pct"], 28)

    def test_blocks_impossible_percentage(self):
        parsed = parse_menu_workbook(build_menu_workbook(fat_value=3.57))

        issues = [item for item in parsed["issues"] if item["severity"] == "error"]
        self.assertTrue(any(item["code"] == "percentage_outlier" for item in issues))


class MenuImportApiTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_db_path = server_app.DB_PATH
        self.old_nutrition_db_path = nutrition_models.DB_PATH
        self.old_upload_dir = server_app.UPLOAD_DIR
        self.old_disable_demo = server_app.DISABLE_DEMO
        db_path = os.path.join(self.temp_dir.name, "student_data.db")
        server_app.DB_PATH = db_path
        nutrition_models.DB_PATH = db_path
        server_app.UPLOAD_DIR = self.temp_dir.name
        server_app.DISABLE_DEMO = False
        server_app._upload_buckets.clear()
        self.client = server_app.app.test_client()
        self.headers = {"X-Demo-Role": "teacher", "X-Demo-Sub": "general"}
        connection = sqlite3.connect(db_path)
        connection.execute(
            """CREATE TABLE students(
                id_card TEXT PRIMARY KEY,
                name TEXT,
                grade_name TEXT,
                class_name TEXT
            )"""
        )
        connection.execute(
            "INSERT INTO students VALUES ('BS001', '示例学生', '三年级', '1班')"
        )
        connection.commit()
        connection.close()

    def tearDown(self):
        server_app._upload_buckets.clear()
        server_app.DB_PATH = self.old_db_path
        nutrition_models.DB_PATH = self.old_nutrition_db_path
        server_app.UPLOAD_DIR = self.old_upload_dir
        server_app.DISABLE_DEMO = self.old_disable_demo
        self.temp_dir.cleanup()

    def _upload(self, workbook_bytes, week=18, parity="even"):
        return self.client.post(
            "/api/menus/upload",
            data={
                "week": str(week),
                "parity": parity,
                "selectionDeadline": "2099-12-31T20:00",
                "image": (io.BytesIO(b"menu image"), "menu.png"),
                "excel": (io.BytesIO(workbook_bytes), "menu.xlsx"),
            },
            content_type="multipart/form-data",
            headers=self.headers,
        )

    def test_draft_is_invisible_until_confirmed_publish(self):
        upload = self._upload(build_menu_workbook(no_service_weekday=3))

        self.assertEqual(upload.status_code, 200)
        draft = upload.get_json()
        self.assertTrue(draft["canPublish"])
        self.assertEqual(draft["preview"]["service_day_count"], 4)
        self.assertEqual(self.client.get("/api/menus").get_json()["menus"], [])

        publish = self.client.post(
            f"/api/menu-imports/{draft['draftId']}/publish",
            headers=self.headers,
        )
        self.assertEqual(publish.status_code, 200)
        self.assertEqual(publish.get_json()["publishedMeals"], 8)

        menu = self.client.get("/api/menus").get_json()["menus"][0]
        self.assertEqual(menu["date_start"], "2026-08-03")
        self.assertEqual(len(menu["service_days"]), 5)
        self.assertEqual(menu["service_days"][2]["service_status"], "no_service")

        connection = sqlite3.connect(server_app.DB_PATH)
        connection.row_factory = sqlite3.Row
        try:
            meals = connection.execute(
                "SELECT * FROM nutrition_meal_plans ORDER BY plan_date, plan_type"
            ).fetchall()
        finally:
            connection.close()
        self.assertEqual(len(meals), 8)
        self.assertEqual(meals[0]["protein_pct"], 24)
        self.assertIsNone(meals[0]["protein_g"])
        self.assertEqual(json.loads(meals[0]["menu_items"]), ["米饭A", "鸡肉时蔬A"])

    def test_teacher_can_fix_blocking_issue_in_draft_without_editing_excel(self):
        upload = self._upload(build_menu_workbook(fat_value=3.57))

        self.assertEqual(upload.status_code, 200)
        draft = upload.get_json()
        self.assertFalse(draft["canPublish"])
        blocked = self.client.post(
            f"/api/menu-imports/{draft['draftId']}/publish",
            headers=self.headers,
        )
        self.assertEqual(blocked.status_code, 409)

        edited_days = draft["preview"]["days"]
        for day in edited_days:
            for meal in day["meals"]:
                meal["fat_pct"] = 28
        saved = self.client.put(
            f"/api/menu-imports/{draft['draftId']}",
            json={
                "title": draft["preview"]["title"],
                "selectionDeadline": "2099-12-31T20:00",
                "days": edited_days,
            },
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.get_json()["canPublish"])
        self.assertEqual(saved.get_json()["issues"], [])

        publish = self.client.post(
            f"/api/menu-imports/{draft['draftId']}/publish",
            headers=self.headers,
        )
        self.assertEqual(publish.status_code, 200)
        self.assertEqual(publish.get_json()["publishedMeals"], 10)

    def test_parent_cannot_edit_menu_draft(self):
        draft = self._upload(build_menu_workbook()).get_json()
        response = self.client.put(
            f"/api/menu-imports/{draft['draftId']}",
            json={"days": draft["preview"]["days"]},
            headers={"X-Demo-Role": "parent", "X-Demo-Kid": "BS001"},
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get("/api/menus").get_json()["menus"], [])

    def test_parent_submits_only_published_service_days(self):
        odd = self._upload(
            build_menu_workbook(start=date(2026, 7, 27), no_service_weekday=3),
            week=17,
            parity="odd",
        ).get_json()
        even = self._upload(
            build_menu_workbook(start=date(2026, 8, 3), no_service_weekday=3),
            week=18,
            parity="even",
        ).get_json()
        for draft in (odd, even):
            response = self.client.post(
                f"/api/menu-imports/{draft['draftId']}/publish",
                headers=self.headers,
            )
            self.assertEqual(response.status_code, 200)

        choices = {str(day): "A" for day in (1, 2, 4, 5)}
        response = self.client.post(
            "/api/meal-choices",
            json={
                "studentIdCard": "BS001",
                "week_odd": 17,
                "week_even": 18,
                "choices": {"odd": choices, "even": choices},
            },
            headers={"X-Demo-Role": "parent", "X-Demo-Kid": "BS001"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["saved"], 8)


if __name__ == "__main__":
    unittest.main()
