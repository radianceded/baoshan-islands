import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from flask import Flask, request, send_file
from openpyxl import Workbook, load_workbook
from PIL import Image
from server import meal_history as h
from server.migrate_meal_history import run


class MealHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.dbs={}
        for campus in ('benbu','baolin','luojing'):
            path=self.root/'student_data.db' if campus=='benbu' else self.root/'campus_data'/campus/'student_data.db'
            path.parent.mkdir(parents=True,exist_ok=True)
            db=sqlite3.connect(path); db.row_factory=sqlite3.Row
            db.executescript('''
                CREATE TABLE weekly_menus(id INTEGER PRIMARY KEY,week_number,parity,date_start,date_end,
                    selection_deadline,service_days_json,notes,default_a_finalized_at);
                CREATE TABLE meal_choices(id_card,name,grade_name,class_name,week_number,parity,weekday,choice,chosen_by,
                    UNIQUE(id_card,week_number,parity,weekday));
                CREATE TABLE nutrition_meal_plans(plan_date,plan_type,plan_name);
            ''')
            db.execute(f'CREATE TABLE students_{campus}(id_card,name,grade_name,class_name,dingtalk_userid)')
            db.executemany(f'INSERT INTO students_{campus} VALUES(?,?,?,?,?)',[
                ('s1','=示例一','三年级','2班','uid1'),('s2','示例二','四年级','1班','uid2')])
            days=[{'weekday':i,'plan_date':f'2026-09-{6+i:02}','service_status':'no_service' if i==5 else 'normal','service_note':''} for i in range(1,8)]
            db.execute('INSERT INTO weekly_menus VALUES(1,2,?,?,?,?,?,?,?)',
                       ('even','2026-09-07','2026-09-13','2026-09-01T00:00',h.pack(days),'','done'))
            db.executemany('INSERT INTO meal_choices VALUES(?,?,?,?,?,?,?,?,?)',[
                (s,n,g,c,2,'even',day,choice,'parent') for s,n,g,c,choice in
                [('s1','=示例一','三年级','2班','A'),('s2','示例二','四年级','1班','B')]
                for day in range(1,8)])
            db.execute("INSERT INTO nutrition_meal_plans VALUES('2026-09-07','A','米饭、青菜')")
            db.commit(); h.migrate(db,'2026s1'); self.dbs[campus]=db
        self.db=self.dbs['benbu']
        app=Flask(__name__); app.config.update(TESTING=True,MEAL_PHOTO_ROOT=str(self.root/'photos'))
        self.user={'role':'teacher','sub_role':'general','campus':'benbu','identity':'test-manager'}
        def book(title,headers,rows):
            wb=Workbook(); ws=wb.active;ws.title=title;ws.append(headers)
            for row in rows:ws.append(row)
            return wb
        def xlsx(wb,name):
            stream=io.BytesIO();wb.save(stream);stream.seek(0)
            return send_file(stream,download_name=name,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        host=SimpleNamespace(app=app,BASE_DIR=self.root,_current_user=lambda:self.user,
            _resolve_campus=lambda:request.args.get('campus','benbu'),get_db=lambda campus:self.dbs[campus],
            _inject_campus_page=lambda name:name,_xlsx_simple_table=book,_xlsx_response=xlsx)
        h.register(host);self.app=app;self.client=app.test_client()

    def tearDown(self):
        for db in self.dbs.values():db.close()
        self.temp.cleanup()

    def archive(self):
        self.db.execute('BEGIN IMMEDIATE')
        ident=h.snapshot(self.db,'benbu',self.db.execute('SELECT * FROM weekly_menus').fetchone())
        self.db.commit();return ident

    def detail(self):return self.client.get('/api/meal-history/2026s1/2').get_json()

    def picture(self,version=0,content=None):
        if content is None:
            stream=io.BytesIO();Image.new('RGB',(24,36),'white').save(stream,format='PNG');content=stream.getvalue()
        return self.client.post('/api/meal-history/2026s1/2/photos',data={
            'version':str(version),'image':(io.BytesIO(content),'image.png')},content_type='multipart/form-data')

    def test_snapshot_counts_weekend_and_excludes_no_service(self):
        self.archive();data=self.detail()
        self.assertEqual(len(data['choices']),12)
        self.assertEqual([(d['A'],d['B']) for d in data['days']],[(1,1)]*4+[(0,0)]+[(1,1)]*2)
        self.assertEqual(data['mealPlans'][0]['plan_name'],'米饭、青菜')

    def test_snapshot_immutable_after_roster_and_choice_changes(self):
        first=self.archive()
        self.db.execute("DELETE FROM students_benbu WHERE id_card='s1'")
        self.db.execute("UPDATE meal_choices SET choice='B' WHERE id_card='s1'");self.db.commit()
        self.assertEqual(self.archive(),first)
        data=self.detail();self.assertTrue(data['changed'])
        self.assertEqual(data['days'][0]['A'],1)
        response=self.client.post('/api/meal-history/2026s1/2/rearchive',json={'version':1,'reason':'核对修正'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.detail()['version'],2)
        self.assertEqual(self.detail()['days'][0]['B'],2)
        old=self.db.execute('SELECT COUNT(*) FROM meal_week_archive_choices WHERE archive_id=?',(first,)).fetchone()[0]
        self.assertEqual(old,12)
        self.user={'role':'parent','campus':'benbu','bound_student_userid':'uid1'}
        self.assertEqual(len(self.detail()['choices']),6)

    def test_parent_and_class_permissions_and_identity_reuse(self):
        self.archive()
        self.user={'role':'parent','campus':'benbu','bound_student_userid':'uid1'}
        self.assertEqual(len(self.detail()['choices']),6)
        self.assertEqual(self.picture().status_code,403)
        self.assertEqual(self.client.get('/api/meal-history/2026s1/2/export.xlsx').status_code,403)
        self.user['bound_student_userid']='new-owner-of-s1'
        self.assertEqual(self.detail()['choices'],[])
        self.user={'role':'teacher','sub_role':'class','campus':'benbu','bound_grade':'三年级','bound_class':'2班'}
        self.assertEqual(len(self.detail()['choices']),6)
        self.db.execute("UPDATE meal_history_settings SET value='2027s1'");self.db.commit()
        self.assertEqual(self.detail()['choices'],[])
        self.assertTrue(self.detail()['detailRestricted'])

    def test_unauthenticated_and_spoofed_campus(self):
        self.archive();self.user['campus']='baolin'
        response=self.client.get('/api/meal-history/2026s1/2?campus=benbu')
        self.assertEqual(response.get_json()['version'],0)
        self.user={'role':'none'}
        self.assertEqual(self.client.get('/api/meal-history').status_code,403)

    def test_read_only_and_open_week_not_archived(self):
        before=self.db.total_changes
        self.client.get('/api/meal-history');self.detail()
        self.assertEqual(self.db.total_changes,before)
        self.db.execute("UPDATE weekly_menus SET selection_deadline='2099-01-01T00:00'");self.db.commit()
        response=self.client.post('/api/meal-history/2026s1/2/rearchive',json={'version':0,'reason':'测试'})
        self.assertEqual(response.status_code,409)
        self.assertIsNone(h.latest(self.db,'2026s1',2))

    def test_photo_version_conflict_keeps_current_and_prior_files(self):
        first=self.picture();self.assertEqual(first.status_code,200)
        ident=first.get_json()['photo']['id']
        self.assertEqual(self.picture().status_code,409)
        self.assertEqual(len(list((self.root/'photos').rglob('*.jpg'))),2)
        self.assertEqual(self.picture(1).status_code,200)
        self.assertEqual(len(list((self.root/'photos').rglob('*.jpg'))),4)
        response=self.client.get(f'/api/meal-photos/{ident}/file');self.assertEqual(response.status_code,200);response.close()
        self.user={'role':'parent','campus':'benbu','bound_student_userid':'uid1'}
        self.assertEqual(self.client.get(f'/api/meal-photos/{ident}/file').status_code,404)
        self.user['campus']='baolin'
        self.assertEqual(self.client.get(f'/api/meal-photos/{ident}/file?campus=benbu').status_code,404)

    def test_invalid_image_and_save_failure_preserve_original(self):
        self.assertEqual(self.picture(content=b'not an image').status_code,400)
        first=self.picture().get_json()['photo']
        with patch('PIL.Image.Image.save',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):self.picture(1,content=self.image_bytes())
        self.assertEqual(self.detail()['photo']['id'],first['id'])
        self.assertEqual(len(list((self.root/'photos').rglob('*.jpg'))),2)

    @staticmethod
    def image_bytes():
        # Fixed PNG fixture independent of Image.save mocking.
        import base64
        return base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')

    def test_export_uses_same_snapshot_and_escapes_formulas(self):
        self.archive();response=self.client.get('/api/meal-history/2026s1/2/export.xlsx')
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb['每日汇总'].cell(2,3).value,self.detail()['days'][0]['A'])
        self.assertEqual(wb['历史选餐明细'].cell(2,3).value,"'=示例一")
        self.assertIn('no-store',response.headers['Cache-Control'])
        response.close()

    def test_revision_conflict_and_migration_idempotence(self):
        report=run(self.root,'benbu','2026s1')
        self.assertEqual(report['archiveWeeks'],[2]);self.assertIsNone(h.latest(self.db,'2026s1',2))
        run(self.root,'benbu','2026s1',True)
        self.assertEqual(self.detail()['source'],'backfill')
        self.assertEqual(run(self.root,'benbu','2026s1',True)['archiveWeeks'],[])
        response=self.client.post('/api/meal-history/2026s1/2/rearchive',json={'version':0,'reason':'过期页面'})
        self.assertEqual(response.status_code,409)
        with self.assertRaises(ValueError):run(self.root,'benbu','2027s1',True)

    def test_storage_is_not_public_and_accel_is_internal(self):
        self.assertEqual(self.client.get('/uploads/meal-photos/benbu/example.jpg').status_code,404)
        self.assertEqual(self.client.get('/uploads/other/../meal-photos/benbu/example.jpg').status_code,404)
        ident=self.picture().get_json()['photo']['id']
        self.app.config['MEAL_PHOTO_X_ACCEL']=True
        response=self.client.get(f'/api/meal-photos/{ident}/file?preview=1')
        self.assertTrue(response.headers['X-Accel-Redirect'].startswith('/_meal_photos/benbu/2026s1/week02/'))

    def test_actual_finalizer_archives_after_default_a_in_same_transaction(self):
        from server.app import _finalize_expired_meal_choices
        self.db.execute('UPDATE weekly_menus SET default_a_finalized_at=NULL')
        self.db.execute("DELETE FROM meal_choices WHERE id_card='s1'")
        self.db.commit()
        result=_finalize_expired_meal_choices(self.db,'benbu')
        self.assertEqual(result['finalized'][0]['added'],6)
        self.assertEqual(self.detail()['days'][0]['B'],1)
        self.assertEqual(self.detail()['days'][0]['A'],1)
        self.assertEqual(self.detail()['source'],'deadline')

    def test_simultaneous_image_decoder_is_rejected(self):
        with h._image_lock:
            self.assertEqual(self.picture().status_code,429)

    def test_backfill_does_not_bind_reused_student_number_to_new_child(self):
        self.db.execute("UPDATE students_benbu SET name='新学生', dingtalk_userid='newuid' WHERE id_card='s1'")
        self.db.commit()
        run(self.root,'benbu','2026s1',True)
        self.user={'role':'parent','campus':'benbu','bound_student_userid':'newuid'}
        self.assertEqual(self.detail()['choices'],[])


if __name__=='__main__':unittest.main()
