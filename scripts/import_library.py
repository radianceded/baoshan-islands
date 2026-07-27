"""导入宝山区实验小学图书馆真实流通数据"""
import openpyxl, sqlite3, re, os, sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, 'student_data.db')
XLSX_PATH = '/root/.openclaw/media/inbound/宝山区实验小学流通数据---273693bc-c114-485a-a540-8164212bdbb2.xlsx'

def clean_name(n):
    """去除姓名中的空格和换行"""
    if not n: return ''
    return str(n).strip().replace('\n','').replace('\r','')

def parse_dept(dept):
    """解析 三年级/（1）班 → (三, 1班)"""
    if not dept: return ('', '')
    dept = str(dept).strip()
    # 匹配: X年级 / （Y）班
    m = re.match(r'(.+?)年级\s*/\s*[（(](\d+)[）)]\s*班', dept)
    if m:
        return (m.group(1).strip(), m.group(2) + '班')
    # 匹配: X年级 Y班
    m = re.match(r'(.+?)年级\s*(\d+)\s*班', dept)
    if m:
        return (m.group(1).strip(), m.group(2) + '班')
    return (dept, '')

def parse_time(t):
    """解析操作时间"""
    if not t: return None
    if isinstance(t, str):
        return t.strip()
    if hasattr(t, 'strftime'):
        return t.strftime('%Y-%m-%d %H:%M:%S')
    return str(t)

def main():
    print('Loading Excel...')
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    print(f'Total rows: {ws.max_row}')

    db = sqlite3.connect(DB_PATH)

    # Create library_records table
    db.execute('DROP TABLE IF EXISTS library_records')
    db.execute('''
        CREATE TABLE library_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            isbn TEXT,
            batch_no TEXT,
            barcode TEXT,
            price REAL,
            action TEXT,
            operation_time TEXT,
            due_date TEXT,
            borrower_name TEXT,
            borrower_account TEXT,
            borrower_dept TEXT,
            borrower_grade TEXT,
            borrower_class TEXT,
            device TEXT,
            student_id_card TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        )
    ''')
    db.execute('CREATE INDEX idx_lib_name ON library_records(borrower_name)')
    db.execute('CREATE INDEX idx_lib_title ON library_records(title)')
    db.execute('CREATE INDEX idx_lib_action ON library_records(action)')
    db.execute('CREATE INDEX idx_lib_time ON library_records(operation_time)')

    # Build student name → id_card map
    student_map = {}
    for r in db.execute('SELECT id_card, name FROM students'):
        student_map[r[1]] = r[0]
    print(f'Student map: {len(student_map)} names')

    inserted = 0
    skipped = 0
    batch = []

    for row_idx in range(2, ws.max_row + 1):  # row 1=title, row 2=header
        row = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]

        title = str(row[0]).strip() if row[0] else ''
        if not title or title == '题名':
            continue

        isbn = str(row[1]).strip() if row[1] else None
        batch_no = str(row[2]).strip() if row[2] else None
        barcode = str(row[3]).strip() if row[3] else None
        price = float(row[4]) if row[4] else None
        action = str(row[5]).strip() if row[5] else None

        # Handle datetime objects
        op_time = parse_time(row[6])
        due_date = parse_time(row[7])

        borrower_name = clean_name(row[8])
        borrower_account = str(row[9]).strip() if row[9] else None
        borrower_dept = str(row[10]).strip() if row[10] else ''

        # Get remaining fields (device, etc.)
        device = str(row[11]).strip() if len(row) > 11 and row[11] else None

        # Parse grade/class from department
        grade, cls = parse_dept(borrower_dept)

        # Map to student id_card
        student_id_card = student_map.get(borrower_name)

        batch.append((
            title, isbn, batch_no, barcode, price, action,
            op_time, due_date, borrower_name, borrower_account,
            borrower_dept, grade, cls, device, student_id_card
        ))

        if len(batch) >= 500:
            db.executemany('''
                INSERT INTO library_records
                (title, isbn, batch_no, barcode, price, action, operation_time,
                 due_date, borrower_name, borrower_account, borrower_dept,
                 borrower_grade, borrower_class, device, student_id_card)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', batch)
            db.commit()
            inserted += len(batch)
            batch = []

        if (row_idx - 2) % 5000 == 0:
            print(f'  Processed {row_idx - 2} rows...')

    # Flush remaining
    if batch:
        db.executemany('''
            INSERT INTO library_records
            (title, isbn, batch_no, barcode, price, action, operation_time,
             due_date, borrower_name, borrower_account, borrower_dept,
             borrower_grade, borrower_class, device, student_id_card)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', batch)
        db.commit()
        inserted += len(batch)

    db.close()
    print(f'\nDone! Inserted: {inserted}')

    # Quick stats
    db2 = sqlite3.connect(DB_PATH)
    print(f'Total records: {db2.execute("SELECT COUNT(*) FROM library_records").fetchone()[0]}')
    print(f'Unique borrowers: {db2.execute("SELECT COUNT(DISTINCT borrower_name) FROM library_records").fetchone()[0]}')
    print(f'Unique books: {db2.execute("SELECT COUNT(DISTINCT title) FROM library_records").fetchone()[0]}')
    print(f'Borrows: {db2.execute("SELECT COUNT(*) FROM library_records WHERE action=\"外借\"").fetchone()[0]}')
    print(f'Returns: {db2.execute("SELECT COUNT(*) FROM library_records WHERE action=\"归还\"").fetchone()[0]}')

    # 陈思麟 stats
    c = db2.execute('SELECT COUNT(*), COUNT(DISTINCT title) FROM library_records WHERE borrower_name="陈思麟"').fetchone()
    print(f'\n陈思麟: {c[0]} records, {c[1]} unique books')
    for r in db2.execute('SELECT title, action, operation_time FROM library_records WHERE borrower_name="陈思麟" ORDER BY operation_time DESC'):
        print(f'  {r[2]} | {r[1]} | {r[0]}')
    db2.close()

if __name__ == '__main__':
    main()
