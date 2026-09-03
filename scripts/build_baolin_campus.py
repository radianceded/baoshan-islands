#!/usr/bin/env python3
"""Build physically isolated campus business/auth databases from read-only rosters."""
from __future__ import annotations

import argparse
import json
import re
import secrets
import sqlite3
import string
from pathlib import Path

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash


PASSWORD_ALPHABET = ''.join(
    ch for ch in string.ascii_letters + string.digits if ch not in '0O1Il'
)
GRADE_MAP = {'一': '一年级', '二': '二年级', '三': '三年级', '四': '四年级', '五': '五年级', '六': '六年级'}
CLASS_RE = re.compile(r'-([一二三四五六])年级(\d+)班(?:\([^)]*\))?$')
TEACHER_CLASS_RE = re.compile(r'^([一二三四五六])(\d+)班?(?:班主任)?$')


def random_password(length: int = 12) -> str:
    return ''.join(secrets.choice(PASSWORD_ALPHABET) for _ in range(length))


def text(value) -> str:
    return str(value or '').strip()


def readonly_connection(path: Path) -> sqlite3.Connection:
    db = sqlite3.connect(f'{path.resolve().as_uri()}?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    return db


def find_header(rows, required: set[str]) -> tuple[tuple[str, ...], object]:
    """Return the first workbook row containing all required headings."""
    for row in rows:
        header = tuple(text(value) for value in row)
        if required.issubset(header):
            return header, rows
    raise SystemExit(f'Cannot find required columns: {sorted(required)}')


def read_students(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['学生信息表1']
    rows = ws.iter_rows(values_only=True)
    header, rows = find_header(rows, {'学生UserID', '学生姓名(必填)', '所在班级(必填)'})
    indexes = {name: header.index(name) for name in ('学生UserID', '学生姓名(必填)', '所在班级(必填)')}
    school_no_index = header.index('学号') if '学号' in header else None
    students = []
    for row in rows:
        user_id = text(row[indexes['学生UserID']])
        name = text(row[indexes['学生姓名(必填)']])
        full_class = text(row[indexes['所在班级(必填)']])
        school_no = text(row[school_no_index]) if school_no_index is not None else ''
        if not user_id and not name:
            continue
        if not user_id or not name or not full_class:
            raise SystemExit(f'Incomplete student row: userId={user_id!r}, name={name!r}')
        match = CLASS_RE.search(full_class)
        if not match:
            raise SystemExit(f'Cannot parse class for {name}: {full_class}')
        students.append({
            'userId': user_id,
            'idCard': user_id,
            'name': name,
            'schoolNo': school_no,
            'grade': GRADE_MAP[match.group(1)],
            'class': f'{match.group(2)}班',
            'fullClass': full_class,
        })
    wb.close()
    user_ids = [student['userId'] for student in students]
    if len(set(user_ids)) != len(user_ids):
        raise SystemExit('Duplicate student UserID in campus roster')
    if not students:
        raise SystemExit('Campus student roster is empty')
    return students


def read_teachers(path: Path) -> tuple[list[dict], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        candidate = tuple(text(value) for value in row)
        if ('员工UserId' in candidate or '员工UserID' in candidate) and '姓名' in candidate:
            header = candidate
            break
    if not header:
        raise SystemExit('Cannot find teacher UserID and name columns')
    user_id_key = '员工UserId' if '员工UserId' in header else '员工UserID'
    mobile_key = '手机号码' if '手机号码' in header else '手机号' if '手机号' in header else None
    position_key = '职位' if '职位' in header else '班主任' if '班主任' in header else None
    indexes = {name: header.index(name) for name in (user_id_key, '姓名')}
    teachers = []
    assignments = []
    for row in rows:
        user_id = text(row[indexes[user_id_key]])
        name = text(row[indexes['姓名']])
        mobile = text(row[header.index(mobile_key)]) if mobile_key else ''
        position = text(row[header.index(position_key)]) if position_key else ''
        if not user_id and not name:
            continue
        if not user_id or not name:
            raise SystemExit(f'Incomplete teacher row: userId={user_id!r}, name={name!r}')
        teacher = {'userId': user_id, 'name': name, 'mobile': mobile, 'position': position}
        teachers.append(teacher)
        match = TEACHER_CLASS_RE.match(position)
        if match:
            assignments.append({
                **teacher,
                'grade': GRADE_MAP[match.group(1)],
                'class': f'{match.group(2)}班',
            })
    wb.close()
    user_ids = [teacher['userId'] for teacher in teachers]
    if len(set(user_ids)) != len(user_ids):
        raise SystemExit('Duplicate teacher UserID in campus roster')
    class_keys = [(assignment['grade'], assignment['class']) for assignment in assignments]
    if len(set(class_keys)) != len(class_keys):
        raise SystemExit('Duplicate class teacher assignment in campus roster')
    return teachers, assignments


def clone_empty_schema(source_path: Path, output_path: Path,
                       skipped_tables: set[str] | None = None) -> sqlite3.Connection:
    source = readonly_connection(source_path)
    output = sqlite3.connect(output_path)
    objects = source.execute(
        '''SELECT type, name, tbl_name, sql FROM sqlite_master
           WHERE sql IS NOT NULL AND name NOT LIKE 'sqlite_%'
           ORDER BY CASE type WHEN 'table' THEN 0 WHEN 'index' THEN 1
                              WHEN 'view' THEN 2 WHEN 'trigger' THEN 3 ELSE 4 END,
                    name'''
    ).fetchall()
    skipped_tables = skipped_tables or set()
    for item in objects:
        if item['name'] in skipped_tables or item['tbl_name'] in skipped_tables:
            continue
        output.execute(item['sql'])
    source.close()
    output.commit()
    return output


def create_business_db(path: Path, schema_source: Path, students: list[dict],
                       teachers: list[dict], assignments: list[dict],
                       campus_id: str, campus_name: str) -> None:
    campus_tables = {
        f'students_{campus_id}',
        f'teacher_roster_{campus_id}',
        f'class_teachers_{campus_id}',
    }
    db = clone_empty_schema(schema_source, path, campus_tables)
    db.execute(f'''CREATE TABLE students_{campus_id}(
        dingtalk_userid TEXT PRIMARY KEY,
        id_card TEXT NOT NULL UNIQUE,
        name TEXT NOT NULL,
        school_no TEXT,
        grade_name TEXT NOT NULL,
        class_name TEXT NOT NULL,
        full_class TEXT,
        student_phone TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        school_name TEXT NOT NULL DEFAULT '',
        gender TEXT
    )''')
    db.execute(f'''CREATE TABLE teacher_roster_{campus_id}(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        dingtalk_userid TEXT NOT NULL UNIQUE,
        name TEXT NOT NULL,
        mobile TEXT,
        department TEXT,
        position TEXT,
        campus TEXT NOT NULL CHECK(campus='{campus_id}'),
        is_admin INTEGER NOT NULL DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )''')
    db.execute(f'''CREATE TABLE class_teachers_{campus_id}(
        class_name TEXT NOT NULL,
        grade_name TEXT NOT NULL,
        teacher_name TEXT NOT NULL,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        dingtalk_userid TEXT NOT NULL,
        PRIMARY KEY (grade_name, class_name)
    )''')
    db.executemany(
        f'''INSERT INTO students_{campus_id}(
               dingtalk_userid, id_card, name, school_no, grade_name,
               class_name, full_class, school_name
           ) VALUES(?,?,?,?,?,?,?,?)''',
        [(s['userId'], s['idCard'], s['name'], s['schoolNo'], s['grade'],
          s['class'], s['fullClass'], campus_name) for s in students],
    )
    db.executemany(
        f'''INSERT INTO teacher_roster_{campus_id}(
               dingtalk_userid, name, mobile, position, campus, is_admin
           ) VALUES(?,?,?,?,?,0)''',
        [(t['userId'], t['name'], t['mobile'], t['position'], campus_id) for t in teachers],
    )
    db.executemany(
        f'''INSERT INTO class_teachers_{campus_id}(
               class_name, grade_name, teacher_name, dingtalk_userid
           ) VALUES(?,?,?,?)''',
        [(a['class'], a['grade'], a['name'], a['userId']) for a in assignments],
    )
    db.commit()
    db.close()


def create_auth_db(path: Path, students: list[dict], teachers: list[dict],
                   assignments: list[dict], general_names: list[str],
                   campus_id: str, campus_name: str) -> dict:
    teachers_by_name = {teacher['name']: teacher for teacher in teachers}
    class_user_ids = {assignment['userId'] for assignment in assignments}
    general_teachers = []
    for name in general_names:
        teacher = teachers_by_name.get(name)
        if not teacher:
            raise SystemExit(f'General teacher missing from teacher roster: {name}')
        if teacher['userId'] in class_user_ids:
            raise SystemExit(f'Role collision: {name} is already a class teacher')
        general_teachers.append(teacher)

    db = sqlite3.connect(path)
    db.execute(f'''CREATE TABLE auth_accounts(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT NOT NULL UNIQUE,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL CHECK(role IN ('parent', 'teacher')),
        sub_role TEXT CHECK(sub_role IS NULL OR sub_role IN ('class', 'general')),
        campus_id TEXT NOT NULL CHECK(campus_id='{campus_id}'),
        display_name TEXT NOT NULL,
        bound_student_userid TEXT,
        bound_id_card TEXT,
        bound_grade TEXT,
        bound_class TEXT,
        is_active INTEGER NOT NULL DEFAULT 1 CHECK(is_active IN (0,1)),
        auth_version INTEGER NOT NULL DEFAULT 1,
        failed_attempts INTEGER NOT NULL DEFAULT 0,
        locked_until INTEGER,
        last_login_at TEXT,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        CHECK(
            (role='parent' AND sub_role IS NULL AND bound_student_userid IS NOT NULL)
            OR (role='teacher' AND sub_role='class' AND bound_grade IS NOT NULL
                AND bound_class IS NOT NULL AND bound_student_userid IS NULL)
            OR (role='teacher' AND sub_role='general' AND bound_student_userid IS NULL
                AND bound_grade IS NULL AND bound_class IS NULL)
        )
    )''')
    db.execute('CREATE INDEX idx_auth_active_role ON auth_accounts(is_active, role, sub_role)')
    credentials = {'campus': campus_id, 'students': [], 'classTeachers': [], 'generalTeachers': []}

    def insert_account(*, username, password, role, sub_role, display_name,
                       student_userid=None, id_card=None, grade=None, class_name=None):
        db.execute(
            '''INSERT INTO auth_accounts(
                   username, password_hash, role, sub_role, campus_id, display_name,
                   bound_student_userid, bound_id_card, bound_grade, bound_class
               ) VALUES(?,?,?,?,?,?,?,?,?,?)''',
            (username, generate_password_hash(password, method='pbkdf2:sha256:100000'),
             role, sub_role, campus_id, display_name, student_userid, id_card, grade, class_name),
        )

    for student in students:
        password = random_password()
        insert_account(
            username=student['userId'], password=password, role='parent', sub_role=None,
            display_name=student['name'], student_userid=student['userId'],
            id_card=student['idCard'], grade=student['grade'], class_name=student['class'],
        )
        credentials['students'].append({
            '校区': campus_name, '年级': student['grade'], '班级': student['class'],
            '姓名': student['name'], '账号': student['userId'], '密码': password,
        })

    for assignment in assignments:
        password = random_password()
        insert_account(
            username=assignment['userId'], password=password, role='teacher',
            sub_role='class', display_name=assignment['name'],
            grade=assignment['grade'], class_name=assignment['class'],
        )
        credentials['classTeachers'].append({
            '校区': campus_name, '年级': assignment['grade'], '班级': assignment['class'],
            '姓名': assignment['name'], '账号': assignment['userId'], '密码': password,
        })

    for teacher in general_teachers:
        password = random_password()
        insert_account(
            username=teacher['userId'], password=password, role='teacher',
            sub_role='general', display_name=teacher['name'],
        )
        credentials['generalTeachers'].append({
            '校区': campus_name, '姓名': teacher['name'],
            '账号': teacher['userId'], '密码': password,
        })

    db.commit()
    db.close()
    return credentials


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--student-roster', type=Path, required=True)
    parser.add_argument('--teacher-roster', type=Path, required=True)
    parser.add_argument('--schema-source-db', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--general-teacher', action='append', default=[])
    parser.add_argument('--campus-id', choices=('baolin', 'luojing'), default='baolin')
    parser.add_argument('--campus-name', default='宝林校区')
    args = parser.parse_args()

    business_path = args.output_dir / 'student_data.db'
    auth_path = args.output_dir / 'auth_accounts.db'
    credentials_path = args.output_dir / 'credentials.json'
    for output in (business_path, auth_path, credentials_path):
        if output.exists():
            raise SystemExit(f'Refusing to overwrite existing output: {output}')
    args.output_dir.mkdir(parents=True, exist_ok=True)

    students = read_students(args.student_roster)
    teachers, assignments = read_teachers(args.teacher_roster)
    create_business_db(
        business_path, args.schema_source_db, students, teachers, assignments,
        args.campus_id, args.campus_name,
    )
    credentials = create_auth_db(
        auth_path, students, teachers, assignments, args.general_teacher,
        args.campus_id, args.campus_name,
    )
    credentials_path.write_text(json.dumps(credentials, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({
        'campus': args.campus_id,
        'students': len(students),
        'teachers': len(teachers),
        'classTeachers': len(assignments),
        'generalTeachers': len(credentials['generalTeachers']),
        'accounts': len(students) + len(assignments) + len(credentials['generalTeachers']),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
